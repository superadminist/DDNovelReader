# -*- coding: utf-8 -*-
"""小说文件解析引擎：把 txt / epub / mobi / azw3 / pdf / docx / html 统一解析为 BookContent。"""
import os
import re
import zipfile
from html.parser import HTMLParser

from .chapterizer import split_chapters, fallback_split

try:
    import docx as _docx
except Exception:  # pragma: no cover
    _docx = None
try:
    import mobi as _mobi
except Exception:  # pragma: no cover
    _mobi = None

SUPPORTED_EXTS = {
    ".txt", ".epub", ".mobi", ".azw3", ".pdf", ".docx", ".html", ".htm", ".zip",
}


class Chapter:
    __slots__ = ("title", "content", "_tts_text", "_tts_map")

    def __init__(self, title, content):
        self.title = (title or "正文").strip()
        self.content = content or ""
        self._tts_text = None
        self._tts_map = None

    def tts_content(self):
        """惰性生成 TTS 朗读用纯净文本与「纯净偏移→原文偏移」映射表。

        显示仍使用原始 content（保留标点/空白/颜文字）；
        朗读使用清理后的纯净文本，避免连续标点/空白/颜文字导致 TTS 静音卡顿。
        """
        if self._tts_text is None:
            from .textproc import preprocess_for_tts

            self._tts_text, self._tts_map = preprocess_for_tts(self.content)
        return self._tts_text, self._tts_map

    def to_dict(self):
        return {"title": self.title, "content": self.content}

    @staticmethod
    def from_dict(d):
        return Chapter(d.get("title", "正文"), d.get("content", ""))


class BookContent:
    def __init__(self, title, author, fmt, chapters):
        self.title = title or "未命名"
        self.author = author or ""
        self.format = fmt or ""
        self.chapters = chapters or []
        # 每章之前的累计字符数，用于 O(1) 计算总进度
        self.cum = [0]
        for c in self.chapters:
            self.cum.append(self.cum[-1] + len(c.content))
        self.total_chars = self.cum[-1]

    def to_dict(self):
        return {
            "v": CONTENT_CACHE_VERSION,
            "title": self.title,
            "author": self.author,
            "format": self.format,
            "chapters": [c.to_dict() for c in self.chapters],
        }

    @staticmethod
    def from_dict(d):
        book = BookContent(
            d.get("title", "未命名"),
            d.get("author", ""),
            d.get("format", ""),
            [Chapter.from_dict(c) for c in d.get("chapters", [])],
        )
        return book


# ---------------- 文本处理 ----------------

def _decode(raw: bytes) -> str:
    """多编码探测解码，保证中文 txt 正确读取。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5", "utf-16"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, ValueError):
            continue
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(raw).best()
        if best is not None and str(best):
            return str(best)
    except Exception:
        pass
    return raw.decode("utf-8", errors="replace")


CONTENT_CACHE_VERSION = 9  # 无章节长文按约 5000 字可靠拆分，旧缓存需重新解析


def normalize_body(text: str) -> str:
    """规范化正文：统一换行、折叠多余空行、段落间以单个换行分隔（不保留空行）。

    空行已移除，段间距完全由阅读区的行距设置控制，避免"每段隔一整行"。
    """
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 网文常见的“全角空格分段”格式：段落之间没有换行，只用 　　 分隔。
    # 若不转换，整章会被连成一段，导致“从该段朗读/高亮跟随”无法定位段落。
    # 先转换双全角空格段首为换行，再处理其余单全角空格（缩进）。
    text = text.replace("\u3000\u3000", "\n")
    text = text.replace("\u3000", "  ")
    text = text.replace("\ufeff", "")
    text = re.sub(r"[ \t]+\n", "\n", text)  # 行尾空格
    text = re.sub(r"\n{2,}", "\n", text)  # 折叠空行：段落间只留一个换行
    text = text.strip("\n")
    if "\n" in text:
        paras = [p.strip("\n").strip() for p in text.split("\n")]
        paras = [p for p in paras if p]
        return "\n".join(paras)
    lines = [ln.rstrip() for ln in text.split("\n") if ln.strip()]
    return "\n".join(lines)


def _make_chapters(text: str) -> list:
    """对整段文本执行章节切分，返回 Chapter 列表。"""
    spl = split_chapters(text)
    if spl:
        return [Chapter(t, normalize_body(b)) for t, b in spl]
    fb = fallback_split(text)
    return [Chapter(t, normalize_body(b)) for t, b in fb]


# ---------------- HTML 提取 ----------------

class _TextExtractor(HTMLParser):
    BLOCK_TAGS = {
        "p", "div", "h1", "h2", "h3", "h4", "h5", "h6",
        "li", "blockquote", "section", "article", "tr", "pre",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._parts = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        if tag in self.BLOCK_TAGS or tag == "br":
            self._parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        if tag in self.BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data):
        if self._skip == 0 and data:
            self._parts.append(data)

    def text(self):
        s = "".join(self._parts)
        return re.sub(r"\n{3,}", "\n\n", s)


def _extract_html(raw: bytes) -> str:
    if isinstance(raw, bytes):
        raw = _decode(raw)
    p = _TextExtractor()
    try:
        p.feed(raw)
    except Exception:
        pass
    return p.text()


# ---------------- 各格式解析 ----------------

def _guess_chapter_title(text: str, max_len=24):
    for line in text.split("\n"):
        line = line.strip(" \t\u3000")
        if line:
            return line[:max_len]
    return None


def _parse_txt(path):
    with open(path, "rb") as f:
        raw = f.read()
    text = _decode(raw)
    title = os.path.splitext(os.path.basename(path))[0]
    return BookContent(title, "", "txt", _make_chapters(text))


def _parse_epub(path):
    zf = zipfile.ZipFile(path)
    names = zf.namelist()
    container = next((n for n in names if n.endswith("META-INF/container.xml")), None)
    if not container:
        raise ValueError("epub 缺少 container.xml，文件可能损坏")
    cdata = zf.read(container).decode("utf-8", "replace")
    m = re.search(r'full-path="([^"]+)"', cdata)
    if not m:
        raise ValueError("epub 的 container.xml 无法解析")
    opf_path = m.group(1)
    opf = zf.read(opf_path).decode("utf-8", "replace")

    title = _first_meta(opf, "dc:title") or os.path.splitext(os.path.basename(path))[0]
    author = _first_meta(opf, "dc:creator") or ""
    author = re.sub(r"<[^>]+>", "", author)

    manifest = {}
    for mm in re.finditer(r"<item\s+([^>]*?)/?>", opf):
        attrs = _tag_attrs(mm.group(1))
        if attrs.get("id"):
            manifest[attrs["id"]] = (attrs.get("href", ""), attrs.get("media-type", ""))

    base = os.path.dirname(opf_path)

    def resolve(href):
        href = href.split("#")[0]
        return os.path.normpath(os.path.join(base, href)).replace("\\", "/")

    spine = []
    for sm in re.finditer(r"<itemref\s+([^>]*?)/?>", opf):
        attrs = _tag_attrs(sm.group(1))
        idref = attrs.get("idref")
        if idref and idref in manifest:
            href, mt = manifest[idref]
            if "html" in mt.lower() or href.lower().endswith((".html", ".xhtml", ".htm")):
                spine.append(resolve(href))

    chapters = []
    seen = set()
    for full in spine:
        if full in seen:
            continue
        seen.add(full)
        try:
            raw = zf.read(full)
        except KeyError:
            continue
        text = normalize_body(_extract_html(raw))
        if len(text) < 20:
            continue
        chapters.append(Chapter(_guess_chapter_title(text) or os.path.basename(full), text))

    if not chapters:
        raise ValueError("epub 未提取到正文内容")
    chapters = _finalize(chapters, "epub")
    return BookContent(title, author, "epub", chapters)


def _parse_mobi(path):
    import glob
    import shutil
    import tempfile

    if _mobi is None:
        raise ValueError("未安装 mobi 解析库，无法解析 mobi/azw3")
    tmp = tempfile.mkdtemp(prefix="novel_reader_mobi_")
    extracted = None
    try:
        title, extracted = _mobi.extract(path)
        html_files = []
        if os.path.isdir(extracted):
            html_files = sorted(glob.glob(os.path.join(extracted, "**", "*.htm*"), recursive=True))
        elif os.path.isfile(extracted):
            html_files = [extracted]
        if not html_files:
            raise ValueError("mobi/azw3 中未找到正文文件")
        chapters = []
        for hf in html_files:
            with open(hf, "rb") as f:
                raw = f.read()
            text = normalize_body(_extract_html(raw))
            if len(text) < 20:
                continue
            chapters.append(Chapter(_guess_chapter_title(text) or os.path.basename(hf), text))
        if not chapters:
            raise ValueError("mobi/azw3 未提取到正文内容")
        chapters = _finalize(chapters, "mobi")
        return BookContent(
            title or os.path.splitext(os.path.basename(path))[0], "", "mobi", chapters
        )
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        if extracted and os.path.isdir(extracted):
            try:
                shutil.rmtree(extracted, ignore_errors=True)
            except Exception:
                pass


def _parse_pdf(path):
    """用 pypdfium2（PDFium，Chrome 同款引擎）解析 PDF 文本。

    比 PyMuPDF 轻量约 10MB+，速度接近，Apache 协议更宽松。
    """
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(path)
    try:
        pages = []
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            pages.append(textpage.get_text_range())
    finally:
        pdf.close()
    text = "\n".join(pages)
    title = os.path.splitext(os.path.basename(path))[0]
    return BookContent(title, "", "pdf", _make_chapters(text))


def _parse_docx(path):
    if _docx is None:
        raise ValueError("未安装 python-docx，无法解析 docx")
    d = _docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    text = "\n".join(lines)
    try:
        title = d.core_properties.title or os.path.splitext(os.path.basename(path))[0]
        author = d.core_properties.author or ""
    except Exception:
        title = os.path.splitext(os.path.basename(path))[0]
        author = ""
    return BookContent(title, author, "docx", _make_chapters(text))


def _parse_html(path):
    with open(path, "rb") as f:
        raw = f.read()
    text = _extract_html(raw)
    title = os.path.splitext(os.path.basename(path))[0]
    return BookContent(title, "", "html", _make_chapters(text))


def _parse_zip(path):
    """解析 zip 压缩包：自动查找其中的 txt / epub 小说文件。"""
    import tempfile
    import shutil

    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise ValueError("zip 文件损坏或不是有效的压缩包")
    except RuntimeError:
        raise ValueError("加密的 zip 暂不支持，请先解压")

    try:
        names = zf.namelist()
        txt_files = [n for n in names
                     if n.lower().endswith(".txt") and not n.endswith("/")
                     and "__MACOSX" not in n]
        epub_files = [n for n in names
                      if n.lower().endswith(".epub") and not n.endswith("/")
                      and "__MACOSX" not in n]

        if not txt_files and not epub_files:
            raise ValueError("压缩包内未找到 txt 或 epub 小说文件")

        if epub_files:
            target = epub_files[0]
            tmp = tempfile.mkdtemp(prefix="ddnovel_zip_")
            try:
                epub_path = os.path.join(tmp, os.path.basename(target))
                with zf.open(target) as src, open(epub_path, "wb") as dst:
                    shutil.copyfileobj(src, dst)
                book = _parse_epub(epub_path)
                book.format = "zip"
                return book
            finally:
                shutil.rmtree(tmp, ignore_errors=True)

        target = max(txt_files, key=lambda n: zf.getinfo(n).file_size)
        raw = zf.read(target)
        text = _decode(raw)
        title = os.path.splitext(os.path.basename(path))[0]
        return BookContent(title, "", "zip", _make_chapters(text))
    finally:
        zf.close()


# ---------------- 辅助 ----------------

def _tag_attrs(s):
    return dict(re.findall(r'([\w:.-]+)="([^"]*)"', s))


def _first_meta(opf, tag):
    m = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", opf, re.S | re.I)
    return m.group(1).strip() if m else ""


def _finalize(chapters, fmt):
    chapters = [c for c in chapters if len(c.content) >= 20]
    if not chapters:
        raise ValueError("未能提取到正文")
    # 单章超大文本：再尝试内部切分（常见于整本书塞进一个 html 的情况）
    if len(chapters) == 1 and len(chapters[0].content) > 40000:
        spl = split_chapters(chapters[0].content)
        if spl and len(spl) >= 2:
            chapters = [Chapter(t, normalize_body(b)) for t, b in spl]
    return chapters


_PARSERS = {
    ".txt": _parse_txt,
    ".epub": _parse_epub,
    ".mobi": _parse_mobi,
    ".azw3": _parse_mobi,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".zip": _parse_zip,
}


def parse_book(path):
    """解析文件为 BookContent。path 必须存在。"""
    ext = os.path.splitext(path)[1].lower()
    if ext not in _PARSERS:
        raise ValueError(f"暂不支持该格式：{ext or '未知'}")
    parser = _PARSERS[ext]
    return parser(path)
